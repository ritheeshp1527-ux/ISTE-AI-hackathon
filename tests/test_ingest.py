"""tests/test_ingest.py

Tests for document extraction and chunking pipeline:
- PDF (PyMuPDF) with 1-based page numbers
- DOCX (python-docx) with headings, tables, units, and page=0
- TXT plain text with page=0
- Chunk word counts (150-250 words) and ~30 word overlap
- Empty files and scanned/empty PDFs (not silently indexed)
- Deterministic chunk IDs
- Title prepending for embedding text while retaining clean passage text
- Full manifest.csv batch ingestion
"""

from pathlib import Path
import docx
import pymupdf as fitz
import pytest

from contracts import (
    CHUNK_13_FIELDS,
    UNKNOWN_PAGE_NUMBER,
    DocumentMetadata,
    EmptyExtractionError,
)
from src.retrieval.ingest import (
    chunk_text,
    extract_document,
    extract_docx,
    extract_pdf,
    extract_txt,
    get_embedding_text,
    ingest_document,
    ingest_manifest,
)


# ============================================================================
# Synthetic Fixtures
# ============================================================================

@pytest.fixture
def synthetic_txt_file(tmp_path: Path) -> Path:
    """Generate a multi-paragraph text file with ~400 words."""
    p1 = " ".join([f"alpha_{i}" for i in range(120)])
    p2 = " ".join([f"beta_{i}" for i in range(140)])
    p3 = " ".join([f"gamma_{i}" for i in range(140)])
    content = f"Introduction to Systems\n\n{p1}\n\n{p2}\n\n{p3}"
    file_path = tmp_path / "system_guide.txt"
    file_path.write_text(content, encoding="utf-8")
    return file_path


@pytest.fixture
def synthetic_docx_file(tmp_path: Path) -> Path:
    """Generate a DOCX file containing headings, paragraphs, and a structured table with units."""
    doc = docx.Document()
    doc.add_heading("Infrastructure Specifications", level=1)
    doc.add_paragraph("This document outlines server hardware requirements and resource limits.")
    
    # Add a structured table with headers, units, and values
    table = doc.add_table(rows=1, cols=4)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Cluster Node"
    hdr_cells[1].text = "Memory (GB)"
    hdr_cells[2].text = "Storage (TB)"
    hdr_cells[3].text = "Max Latency (ms)"

    row_data = [
        ("node-primary-01", "256 GB", "10 TB", "12 ms"),
        ("node-replica-02", "128 GB", "5 TB", "25 ms"),
        ("node-analytics-03", "512 GB", "20 TB", "45 ms"),
    ]
    for row in row_data:
        r_cells = table.add_row().cells
        for col_idx, text in enumerate(row):
            r_cells[col_idx].text = text

    # Add paragraphs of text
    p_content = " ".join([f"spec_word_{i}" for i in range(160)])
    doc.add_paragraph(p_content)

    file_path = tmp_path / "infrastructure.docx"
    doc.save(str(file_path))
    return file_path


@pytest.fixture
def synthetic_pdf_file(tmp_path: Path) -> Path:
    """Generate a 3-page PDF with distinct text on each page using PyMuPDF."""
    doc = fitz.open()

    # Page 1
    p1 = doc.new_page()
    text_p1 = "Page 1 Content: Executive Summary.\n\n" + " ".join([f"summary_word_{i}" for i in range(160)])
    p1.insert_text((50, 72), text_p1)

    # Page 2
    p2 = doc.new_page()
    text_p2 = "Page 2 Content: Financial Analysis.\n\n" + " ".join([f"finance_word_{i}" for i in range(170)])
    p2.insert_text((50, 72), text_p2)

    # Page 3
    p3 = doc.new_page()
    text_p3 = "Page 3 Content: Risk Assessment.\n\n" + " ".join([f"risk_word_{i}" for i in range(80)])
    p3.insert_text((50, 72), text_p3)

    file_path = tmp_path / "annual_report.pdf"
    doc.save(str(file_path))
    doc.close()
    return file_path


@pytest.fixture
def empty_txt_file(tmp_path: Path) -> Path:
    f = tmp_path / "empty.txt"
    f.write_text("", encoding="utf-8")
    return f


@pytest.fixture
def empty_pdf_file(tmp_path: Path) -> Path:
    """A PDF with a page containing no text (simulating an un-OCRed scanned PDF)."""
    doc = fitz.open()
    doc.new_page()  # Blank page
    file_path = tmp_path / "blank_scanned.pdf"
    doc.save(str(file_path))
    doc.close()
    return file_path


# ============================================================================
# Tests: Extraction and Metadata
# ============================================================================

def test_extract_txt_and_ingest(synthetic_txt_file: Path):
    meta = DocumentMetadata(
        doc_id="DOC-TXT-01",
        title="System Guide",
        source_path=str(synthetic_txt_file),
        source_type="txt",
        version="1.0",
        status="active",
        access_group="engineering",
        modified_at="2026-10-01",
        supersedes=None,
    )
    chunks = ingest_document(meta)
    assert len(chunks) >= 2, f"Expected multiple chunks, got {len(chunks)}"

    for chunk in chunks:
        # Verify all 13 fields
        for field in CHUNK_13_FIELDS:
            assert getattr(chunk, field) is not None or field == "supersedes"

        assert chunk.doc_id == "DOC-TXT-01"
        assert chunk.title == "System Guide"
        assert chunk.access_group == "engineering"
        assert chunk.version == "1.0"
        # Unknown page convention for TXT: page == 0
        assert chunk.page == UNKNOWN_PAGE_NUMBER
        assert chunk.page == 0

        # Clean passage text
        assert not chunk.text.startswith("System Guide\n\n")

        # Embedding text prepends title
        emb_text = get_embedding_text(chunk)
        assert emb_text.startswith("System Guide\n\n")
        assert chunk.embedding_text.startswith("System Guide\n\n")


def test_extract_docx_and_ingest(synthetic_docx_file: Path):
    meta = DocumentMetadata(
        doc_id="DOC-DOCX-01",
        title="Infrastructure Specs",
        source_path=str(synthetic_docx_file),
        source_type="docx",
        version="2.0",
        status="active",
        access_group="devops",
        modified_at="2026-10-05",
        supersedes="DOC-DOCX-00",
    )
    chunks = ingest_document(meta)
    assert len(chunks) >= 1

    # Check that table and headings were extracted
    all_text = " ".join(c.text for c in chunks)
    assert "# Infrastructure Specifications" in all_text
    assert "node-primary-01" in all_text
    assert "256 GB" in all_text
    assert "12 ms" in all_text

    for chunk in chunks:
        assert chunk.doc_id == "DOC-DOCX-01"
        assert chunk.access_group == "devops"
        assert chunk.version == "2.0"
        assert chunk.supersedes == "DOC-DOCX-00"
        # Unknown page convention for DOCX: page == 0
        assert chunk.page == UNKNOWN_PAGE_NUMBER
        assert chunk.page == 0


def test_extract_pdf_preserves_1_based_page_numbers(synthetic_pdf_file: Path):
    meta = DocumentMetadata(
        doc_id="DOC-PDF-01",
        title="Annual Report",
        source_path=str(synthetic_pdf_file),
        source_type="pdf",
        version="1.0",
        status="active",
        access_group="executive",
        modified_at="2026-10-10",
        supersedes=None,
    )
    chunks = ingest_document(meta)
    assert len(chunks) >= 3

    # Verify that PDF pages are strictly 1-based (1, 2, 3)
    pages_found = {chunk.page for chunk in chunks}
    assert 1 in pages_found
    assert 2 in pages_found
    assert 3 in pages_found
    assert 0 not in pages_found  # 0 must NOT be used for valid PDF pages

    # Verify content corresponds to correct page number
    for chunk in chunks:
        if chunk.page == 1:
            assert "Page 1 Content" in chunk.text
        elif chunk.page == 2:
            assert "Page 2 Content" in chunk.text
        elif chunk.page == 3:
            assert "Page 3 Content" in chunk.text


# ============================================================================
# Tests: Chunking Size & Overlap
# ============================================================================

def test_chunking_word_counts_and_overlap():
    """Verify chunks are approximately 150-250 words with ~30 words of overlap."""
    paragraphs = [
        " ".join([f"word_a_{i}" for i in range(160)]),
        " ".join([f"word_b_{i}" for i in range(160)]),
        " ".join([f"word_c_{i}" for i in range(160)]),
    ]
    sample_text = "\n\n".join(paragraphs)

    chunks = chunk_text(sample_text, target_min_words=150, target_max_words=250, overlap_words=30)
    assert len(chunks) >= 2

    for c in chunks:
        words = c.split()
        # Chunks should be approximately 150-250 words
        assert len(words) >= 120, f"Chunk too small: {len(words)} words"
        assert len(words) <= 300, f"Chunk too large: {len(words)} words"

    # Check overlap between consecutive chunks:
    # The end of chunk[0] should share ~30 words with the beginning of chunk[1]
    c0_words = chunks[0].split()
    c1_words = chunks[1].split()
    overlap_set_0 = set(c0_words[-30:])
    overlap_set_1 = set(c1_words[:35])
    common_words = overlap_set_0.intersection(overlap_set_1)
    assert len(common_words) >= 15, f"Expected ~30 overlap words, found {len(common_words)}"


# ============================================================================
# Tests: Empty Files & Scanned PDFs
# ============================================================================

def test_empty_files_not_silently_indexed(empty_txt_file: Path, empty_pdf_file: Path):
    """Empty files and un-OCRed scanned PDFs must not silently produce empty chunks."""
    meta_txt = DocumentMetadata(
        doc_id="DOC-EMPTY-TXT",
        title="Empty TXT",
        source_path=str(empty_txt_file),
        source_type="txt",
        modified_at="2026-01-01",
    )
    # Default non-raising mode returns empty list (no empty text indexed)
    chunks_txt = ingest_document(meta_txt, raise_on_empty=False)
    assert len(chunks_txt) == 0

    # Raising mode clearly raises EmptyExtractionError
    with pytest.raises(EmptyExtractionError, match="Empty text extracted"):
        ingest_document(meta_txt, raise_on_empty=True)

    meta_pdf = DocumentMetadata(
        doc_id="DOC-BLANK-PDF",
        title="Blank Scanned PDF",
        source_path=str(empty_pdf_file),
        source_type="pdf",
        modified_at="2026-01-01",
    )
    chunks_pdf = ingest_document(meta_pdf, raise_on_empty=False)
    assert len(chunks_pdf) == 0

    with pytest.raises(EmptyExtractionError, match="scanned PDF requiring OCR"):
        ingest_document(meta_pdf, raise_on_empty=True)


# ============================================================================
# Tests: Deterministic Chunk IDs & Re-ingestion
# ============================================================================

def test_deterministic_chunk_ids_on_reingestion(synthetic_txt_file: Path):
    meta = DocumentMetadata(
        doc_id="DOC-STABLE-01",
        title="System Guide",
        source_path=str(synthetic_txt_file),
        source_type="txt",
        version="1.0",
        modified_at="2026-10-01",
    )
    run1 = ingest_document(meta)
    run2 = ingest_document(meta)

    assert len(run1) == len(run2)
    for c1, c2 in zip(run1, run2):
        assert c1.chunk_id == c2.chunk_id
        assert c1.text == c2.text
        assert c1.chunk_index == c2.chunk_index

    # Different version produces different chunk IDs
    meta_v2 = meta.model_copy(update={"version": "2.0"})
    run_v2 = ingest_document(meta_v2)
    assert run1[0].chunk_id != run_v2[0].chunk_id


# ============================================================================
# Tests: Full Manifest CSV Batch Ingestion
# ============================================================================

def test_ingest_manifest_batch(
    tmp_path: Path,
    synthetic_txt_file: Path,
    synthetic_docx_file: Path,
    synthetic_pdf_file: Path,
):
    manifest_csv = tmp_path / "manifest.csv"
    manifest_csv.write_text(
        "doc_id,title,source_path,source_type,version,status,access_group,modified_at,supersedes\n"
        f"DOC-TXT,System Guide,{synthetic_txt_file.name},txt,1.0,active,public,2026-10-01,\n"
        f"DOC-DOCX,Infra Specs,{synthetic_docx_file.name},docx,1.1,active,devops,2026-10-02,\n"
        f"DOC-PDF,Annual Report,{synthetic_pdf_file.name},pdf,1.0,active,finance,2026-10-03,\n",
        encoding="utf-8",
    )

    all_chunks = ingest_manifest(manifest_csv, base_dir=tmp_path)
    assert len(all_chunks) >= 6

    # Verify representation of all three document types
    types_found = {c.source_type for c in all_chunks}
    assert types_found == {"txt", "docx", "pdf"}

    # Verify access groups are assigned correctly
    groups_found = {c.access_group for c in all_chunks}
    assert "public" in groups_found
    assert "devops" in groups_found
    assert "finance" in groups_found
